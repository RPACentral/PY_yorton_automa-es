import pandas as pd
import pyodbc
from functions.data_base import db_connection # função para conectar com banco de dados
from functions.colors import green, red # função para utilizar cores nos prints
from dateutil import parser
from datetime import datetime, timedelta
from openpyxl.styles import NamedStyle
import os

# Função para obter o primeiro e último dia do mês
def get_month_range(year, month):
    first_day = datetime(year, month, 1)
    last_day = (datetime(year, month + 1, 1) - timedelta(days=1)) if month < 12 else datetime(year + 1, 1, 1) - timedelta(days=1)
    return first_day, last_day

# Obter o mês e ano atuais
date = datetime.now()
year = date.year
month = date.month

# Obter o intervalo de datas para o mês atual
month_start, month_end = get_month_range(year, month)

# Ler a planilha .xlsm
input_file = 'H:/Tecnologia/EQUIPE - DADOS/0 - Microsoft Power B.I/BASES PARA CARGA/RELACIONAMENTO_MEDICO/BASE FECHAMENTO OUVIDORIA.xlsm'
sheet_name = 'BASE'

df = pd.read_excel(input_file, sheet_name=sheet_name)

# Filtrar os dados para o mês atual
filtered_df = df[(df['mês ref'] >= month_start) & (df['mês ref'] <= month_end)].copy()

# Renomear as colunas para o formato desejado
column_mapping = {
    'TASK NAME': 'NOME_TAREFA',
    'ASSIGNEE': 'RESPONSAVEL',
    'STATUS': 'STATUS',
    'DATE CREATED': 'DATA_CRIACAO',
    'DATE CLOSED': 'DATA_FECHAMENTO',
    'SLA': 'SLA',
    'mês ref': 'MES_REF',
    'SLA Ajustado': 'SLA_AJUSTADO',
    'Filtro': 'FILTRO'
}

# Aplicar a renomeação das colunas
filtered_df.rename(columns=column_mapping, inplace=True)

def parse_relative_date(date_str):
    # Obtém a data e hora atuais
    now = datetime.now()
    
    # Verifica se a entrada é uma string
    if isinstance(date_str, str):
        # Converte a string para minúsculas e remove espaços extras
        date_str = date_str.lower().strip()
        
        # Verifica se a string representa o dia de hoje
        if date_str in ['hoje', 'today']:
            # Retorna a data atual no formato DD/MM/AAAA
            return now.strftime('%d/%m/%Y')
        
        # Verifica se a string representa o dia anterior
        elif date_str in ['ontem', 'yesterday']:
            # Retorna a data de ontem no formato DD/MM/AAAA
            return (now - timedelta(days=1)).strftime('%d/%m/%Y')
        
        # Verifica se a string contém a expressão 'dias atrás' ou 'days ago'
        elif 'dias atrás' in date_str or 'days ago' in date_str:
            try:
                # Extrai o número de dias da string e converte para inteiro
                days = int(date_str.split()[0])
                # Retorna a data calculada (hoje menos o número de dias) no formato DD/MM/AAAA
                return (now - timedelta(days=days)).strftime('%d/%m/%Y')
            except:
                print(red("Erro ao converter dados!(x dias atrás)"))
                exit()
        
        else:
            # Tenta analisar a string de data em um formato padrão
            try:
                # Usa a função 'parse' para converter a string para um objeto datetime
                return parser.parse(date_str, fuzzy=True).strftime('%d/%m/%Y')
            except:
                print(red("Erro ao converter string para data!"))
                exit()
    
    # Se a entrada não for uma string, cancela o processo
    print(red("Erro ao ler dados com string!"))
    exit()


# Aplicar a conversão de datas relativas para as colunas especificadas
if 'DATA_CRIACAO' in filtered_df.columns:
    filtered_df['DATA_CRIACAO'] = filtered_df['DATA_CRIACAO'].apply(parse_relative_date)
if 'DATA_FECHAMENTO' in filtered_df.columns:
    filtered_df['DATA_FECHAMENTO'] = filtered_df['DATA_FECHAMENTO'].apply(parse_relative_date)

# Criar um estilo de célula para data
date_style = NamedStyle(name='date_style', number_format='DD/MM/YYYY')

# Definir o caminho fixo onde salvar o arquivo
output_directory = 'C:/Users/yorton.filho/Downloads/'
output_filename = 'OUVIDORIA_FORMATADA.xlsx'
output_file = os.path.join(output_directory, output_filename)

# Escrever os dados filtrados na nova planilha e aplicar formatação
with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
    # Converter DataFrame para Excel
    filtered_df.to_excel(writer, index=False, sheet_name='Dados Filtrados')
    
    # Obter o workbook e a planilha
    workbook = writer.book
    sheet = workbook['Dados Filtrados']
    
    # Aplicar a formatação de data para as colunas especificadas
    for col in ['DATA_CRIACAO', 'DATA_FECHAMENTO', 'MES_REF']:
        col_index = filtered_df.columns.get_loc(col) + 1  # Obter o índice da coluna
        for row in range(2, sheet.max_row + 1):  # Começa na linha 2 para pular o cabeçalho
            cell = sheet.cell(row=row, column=col_index)
            if isinstance(cell.value, str):
                try:
                    # Converte a string para um objeto datetime se estiver no formato esperado
                    cell.value = datetime.strptime(cell.value, '%d/%m/%Y')
                except:
                    print(red("Erro ao converter string para data!"))
            if isinstance(cell.value, datetime):
                cell.style = date_style

print(green(f"Dados filtrados foram salvos em {output_file}"))

# ---------------------------------------------------------------- IMPORTAÇÃO PARA O BANCO

try:
    # Ler o arquivo .xlsx com pandas, usando a primeira linha como cabeçalho
    data = pd.read_excel(output_file, engine='openpyxl', header=0)

    # Imprimir as primeiras linhas e colunas do DataFrame para depuração
    # print("Primeiras linhas do DataFrame:")
    # print(data.head())
    # print("\nNomes das colunas do DataFrame:")
    # print(data.columns.tolist())

    try:
        with db_connection() as connection:  # conectando no banco
            with connection.cursor() as cursor:  # abrindo central para query
                table = 'DADOS_OUVIDORIA'
                
                # Obter a estrutura da tabela Oracle
                cursor.execute(f"SELECT COLUMN_NAME FROM ALL_TAB_COLUMNS WHERE TABLE_NAME = '{table}'")
                columns = [row[0] for row in cursor.fetchall()]
                num_columns = len(columns)
                
                # print(f"Número de colunas na tabela Oracle: {num_columns}")
                # print(f"Nomes das colunas na tabela Oracle: {columns}")

                # Verifica se o número de colunas na tabela Oracle corresponde ao número de colunas no DataFrame
                if len(data.columns) != num_columns:
                    print(red(f"Número de colunas no DataFrame ({len(data.columns)}) não corresponde ao número de colunas na tabela Oracle ({num_columns})."))
                    exit()
                
                # Reordenar as colunas do DataFrame para corresponder à ordem da tabela Oracle
                data = data[columns]
                
                # Converter os tipos de dados, se necessário
                for col in columns:
                    if 'DATA' in col.upper():
                        # As datas já foram convertidas acima
                        continue
                    elif 'NUM' in col.upper():
                        # Converte colunas numéricas para o formato numérico
                        data[col] = pd.to_numeric(data[col], errors='coerce')

                # Query para limpar a tabela
                date = datetime.now()
                month = f'{date.month:02d}'
                year = date.year
                delete_command = f"DELETE FROM {table} WHERE MES_REF = '01/{month}/{year}'"
                
                # Deletar dados do banco
                try:
                    cursor.execute(delete_command)
                    print(green("Dados deletados com sucesso!"))
                except pyodbc.Error as e:
                    print(red(f"Erro ao deletar dados: {e}"))
                    exit()

                # Criar a query de inserção
                placeholders = ', '.join(['?' for _ in range(num_columns)])
                insert_command = f"INSERT INTO {table} ({', '.join(columns)}) VALUES ({placeholders})"
                
                # Inserir dados na tabela Oracle
                for index, row in data.iterrows():
                    values = row.tolist()  # Converte a linha em lista de valores
                    
                    # Substituir strings vazias por None para garantir que sejam tratados como NULL
                    values = [None if v == '' or pd.isna(v) else v for v in values]

                    # Inserindo os dados no banco
                    try:
                        cursor.execute(insert_command, values)
                    except pyodbc.Error as e:
                        print(red(f"Erro ao inserir dados: {e}"))
                        exit()

                connection.commit()
                print(green("Dados importados com sucesso!"))

    except pyodbc.Error as e:
        print(red(f"Erro ao conectar ou interagir com o banco de dados: {e}"))

except FileNotFoundError as e:
    print(red(f"Erro: {e}"))
